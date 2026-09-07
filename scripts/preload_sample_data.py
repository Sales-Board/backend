#!/usr/bin/env python3
from __future__ import annotations

import argparse
import re
import sys
from datetime import UTC, datetime, timedelta
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import func
from sqlalchemy import select
from sqlalchemy.orm import Session

ROOT_DIR = Path(__file__).resolve().parents[1]
if str(ROOT_DIR) not in sys.path:
    sys.path.insert(0, str(ROOT_DIR))

from app.db.database import get_session_factory
from app.models.call import Call
from app.models.campaign import Campaign
from app.models.customer import Customer
from app.models.engagement_event import EngagementEvent
from app.models.followup import Followup
from app.models.lead import Lead
from app.models.lead_assignment import LeadAssignment
from app.models.lead_outcome import LeadOutcome
from app.models.lead_timeline_event import LeadTimelineEvent
from app.models.product import Product
from app.models.task import Task
from app.models.website_event import WebsiteEvent

DATASET_PATH = ROOT_DIR / "sample_data.xlsx"


def _slug(value: str, max_len: int = 40) -> str:
    normalized = re.sub(r"[^a-z0-9]+", "-", value.lower()).strip("-")
    return normalized[:max_len] or "x"


def _str(value: object | None, max_len: int | None = None) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    if max_len is not None:
        return text[:max_len]
    return text


def _as_int(value: object | None, default: int = 0) -> int:
    if value is None:
        return default
    try:
        return int(float(str(value)))
    except (TypeError, ValueError):
        return default


def _as_float(value: object | None, default: float = 0.0) -> float:
    if value is None:
        return default
    try:
        return float(str(value))
    except (TypeError, ValueError):
        return default


def _as_bool(value: object | None) -> bool:
    text = _str(value).lower()
    return text in {"1", "true", "t", "yes", "y", "repeat"}


def _source_channel(row: dict[str, object]) -> str:
    for key in ("CRM_Channel", "CRM_Source", "WEB_Last_Touch_Channel", "CRM_Data_Source_Platform"):
        value = _str(row.get(key), 64)
        if value:
            return value
    return "unknown"


def _source_medium(row: dict[str, object]) -> str:
    for key in ("CRM_Data_Medium", "CRM_Lead_Source", "CRM_Lead_Type", "WEB_Device_Type"):
        value = _str(row.get(key), 64)
        if value:
            return value
    return "unknown"


def _normalize_status(row: dict[str, object]) -> str:
    value = " ".join(
        [
            _str(row.get("Label_Source_Lead_Status")),
            _str(row.get("Label_Source_Disposition")),
        ]
    ).lower()
    if any(token in value for token in ("converted", "payment", "closed", "sale")):
        return "converted"
    if any(token in value for token in ("interested", "follow", "contactable", "callback")):
        return "qualified"
    if any(token in value for token in ("lost", "invalid", "not interested", "reject")):
        return "lost"
    return "new"


def _compute_lead_score(row: dict[str, object]) -> float:
    validity = _as_float(row.get("LABEL_Customer_Validity"), default=1.0)
    connect_rate = _as_float(row.get("CDR_Connect_Rate"), default=0.0)
    page_views = _as_int(row.get("Page_Views"), default=0)
    msg_engaged = _as_int(row.get("MSG_Engaged"), default=0)

    score = (max(1.0, min(validity, 3.0)) / 3.0) * 40.0
    score += max(0.0, min(connect_rate, 1.0)) * 30.0
    score += min(page_views, 30) * 0.8
    score += min(msg_engaged, 20) * 0.6
    return round(min(score, 100.0), 2)


def _priority_from_score(score: float) -> str:
    if score >= 70:
        return "high"
    if score >= 40:
        return "medium"
    return "low"


def _stage_from_status(status: str) -> str:
    if status == "converted":
        return "completed"
    if status == "qualified":
        return "nurturing"
    if status == "lost":
        return "closed"
    return "generated"


def _next_action_hint(row: dict[str, object]) -> str:
    text = " ".join(
        [
            _str(row.get("Label_Source_Disposition")),
            _str(row.get("CDR_Top_Q_Type")),
            _str(row.get("Label_Source_Lead_Status")),
        ]
    ).lower()
    if any(token in text for token in ("callback", "follow", "call")):
        return "call"
    if any(token in text for token in ("payment", "proposal", "otp", "convert")):
        return "close"
    return "engage"


def _load_excel_rows(path: Path) -> tuple[list[str], list[dict[str, object]]]:
    if not path.exists():
        raise FileNotFoundError(f"Sample dataset not found at: {path}")

    workbook = load_workbook(path, read_only=True, data_only=True)
    if "dataset" not in workbook.sheetnames:
        raise ValueError("sample_data.xlsx is missing required 'dataset' sheet")

    sheet = workbook["dataset"]
    iterator = sheet.iter_rows(values_only=True)
    try:
        header_row = next(iterator)
    except StopIteration as exc:
        raise ValueError("dataset sheet is empty") from exc

    headers = [(_str(name) or f"column_{idx + 1}") for idx, name in enumerate(header_row)]

    rows: list[dict[str, object]] = []
    for values in iterator:
        if values is None:
            continue
        if all(value is None or _str(value) == "" for value in values):
            continue
        row = {headers[idx]: values[idx] if idx < len(values) else None for idx in range(len(headers))}
        rows.append(row)

    return headers, rows


def _reset_domain_tables(db: Session) -> None:
    from app.models.ai_prediction import AIPrediction
    from app.models.decision_recommendation import DecisionRecommendation

    for model in [
        Followup,
        Task,
        LeadOutcome,
        LeadTimelineEvent,
        LeadAssignment,
        DecisionRecommendation,
        AIPrediction,
        Call,
        WebsiteEvent,
        EngagementEvent,
        Lead,
        Campaign,
        Product,
        Customer,
    ]:
        db.query(model).delete()
    db.flush()


def _remove_smoke_records(db: Session) -> None:
    db.query(Customer).filter(Customer.external_customer_id.like("SMOKE-%")).delete(synchronize_session=False)
    db.query(Customer).filter(Customer.external_customer_id.like("SAMPLE-%")).delete(synchronize_session=False)
    db.query(Campaign).filter(Campaign.code.like("SMOKE-%")).delete(synchronize_session=False)
    db.query(Campaign).filter(Campaign.code.like("SAMPLE-%")).delete(synchronize_session=False)
    db.query(Product).filter(Product.code.like("SMOKE-%")).delete(synchronize_session=False)
    db.query(Product).filter(Product.code.like("SAMPLE-%")).delete(synchronize_session=False)
    db.flush()


def _count(db: Session, model, *filters) -> int:
    stmt = select(func.count()).select_from(model)
    for rule in filters:
        stmt = stmt.where(rule)
    return int(db.scalar(stmt) or 0)


def main() -> None:
    parser = argparse.ArgumentParser(description="Load real sample_data.xlsx records into PostgreSQL")
    parser.add_argument("--reset", action="store_true", help="Delete existing domain rows before load")
    parser.add_argument("--keep-smoke", action="store_true", help="Do not remove existing SMOKE/SAMPLE rows")
    parser.add_argument(
        "--dataset-path",
        default=str(DATASET_PATH),
        help="Path to source sample_data.xlsx (default: backend/sample_data.xlsx)",
    )
    args = parser.parse_args()

    session_factory = get_session_factory()
    if session_factory is None:
        raise RuntimeError("Database is not configured")

    dataset_path = Path(args.dataset_path)
    columns, rows = _load_excel_rows(dataset_path)

    created = {
        "products": 0,
        "campaigns": 0,
        "customers": 0,
        "leads": 0,
        "lead_assignments": 0,
        "lead_timeline_events": 0,
        "tasks": 0,
        "followups": 0,
        "calls": 0,
        "engagement_events": 0,
        "website_events": 0,
    }

    with session_factory() as db:
        if args.reset:
            _reset_domain_tables(db)
        elif not args.keep_smoke:
            _remove_smoke_records(db)

        for row in rows:
            customer_external_id = _str(row.get("Customer_ID"), 64)
            if not customer_external_id:
                continue

            product_code = _str(row.get("CRM_Product_Code"), 64) or f"PRD-{_slug(_str(row.get('CRM_Product_Name')) or customer_external_id, 50)}"
            product_name = _str(row.get("CRM_Product_Name"), 150) or _str(row.get("WEB_Plan_Type"), 150) or "Unknown Product"
            product_category = _str(row.get("WEB_Plan_Variant"), 100) or _str(row.get("CRM_Lead_Type"), 100) or None
            product = db.scalar(select(Product).where(Product.code == product_code))
            if product is None:
                product = Product(
                    code=product_code,
                    name=product_name,
                    category=product_category,
                    is_active=True,
                )
                db.add(product)
                db.flush()
                created["products"] += 1

            campaign_seed = _str(row.get("CRM_UTM_Campaign"), 64) or _str(row.get("CRM_Data_Source_Platform"), 64) or _str(row.get("CRM_Channel"), 64)
            campaign_code = _str(campaign_seed, 64) or f"CMP-{_slug(product_code, 50)}"
            campaign_name = _str(row.get("CRM_UTM_Campaign"), 150) or _str(row.get("CRM_Data_Source_Platform"), 150) or f"{product_name} Campaign"
            campaign_channel = _source_channel(row)
            campaign = db.scalar(select(Campaign).where(Campaign.code == campaign_code))
            if campaign is None:
                campaign = Campaign(
                    code=campaign_code,
                    name=campaign_name,
                    channel=campaign_channel,
                    status="active",
                )
                db.add(campaign)
                db.flush()
                created["campaigns"] += 1

            customer = db.scalar(select(Customer).where(Customer.external_customer_id == customer_external_id))
            if customer is None:
                customer = Customer(
                    external_customer_id=customer_external_id,
                    first_name=None,
                    last_name=None,
                    crm_gender=_str(row.get("CRM_Gender"), 32) or None,
                    crm_age_band=_str(row.get("CRM_Age_Band"), 32) or None,
                    crm_income_band=_str(row.get("CRM_Income_Band"), 32) or None,
                    crm_occupation=_str(row.get("CRM_Occupation"), 100) or None,
                    crm_education=_str(row.get("CRM_Education"), 100) or None,
                    crm_tobacco_user=_str(row.get("CRM_Tobacco_User"), 16) or None,
                    crm_nonresident_flag=_str(row.get("CRM_NonResident_Flag"), 16) or None,
                    crm_existing_plan_flag=_str(row.get("CRM_Existing_Plan_Flag"), 64) or None,
                )
                db.add(customer)
                db.flush()
                created["customers"] += 1
            else:
                # Backfill CRM columns for existing rows if missing.
                customer.crm_gender = customer.crm_gender or (_str(row.get("CRM_Gender"), 32) or None)
                customer.crm_age_band = customer.crm_age_band or (_str(row.get("CRM_Age_Band"), 32) or None)
                customer.crm_income_band = customer.crm_income_band or (_str(row.get("CRM_Income_Band"), 32) or None)
                customer.crm_occupation = customer.crm_occupation or (_str(row.get("CRM_Occupation"), 100) or None)
                customer.crm_education = customer.crm_education or (_str(row.get("CRM_Education"), 100) or None)
                customer.crm_tobacco_user = customer.crm_tobacco_user or (_str(row.get("CRM_Tobacco_User"), 16) or None)
                customer.crm_nonresident_flag = customer.crm_nonresident_flag or (_str(row.get("CRM_NonResident_Flag"), 16) or None)
                customer.crm_existing_plan_flag = customer.crm_existing_plan_flag or (
                    _str(row.get("CRM_Existing_Plan_Flag"), 64) or None
                )

            source_channel = _source_channel(row)
            source_medium = _source_medium(row)
            status_value = _normalize_status(row)
            lead_score = _compute_lead_score(row)
            priority = _priority_from_score(lead_score)
            current_stage = _stage_from_status(status_value)
            current_section = _str(row.get("CRM_Department"), 64) or "intake"
            current_handler = _str(row.get("CRM_Source"), 100) or None

            lead = db.scalar(
                select(Lead).where(
                    Lead.customer_id == customer.id,
                    Lead.campaign_id == campaign.id,
                )
            )
            if lead is None:
                lead = Lead(
                    customer_id=customer.id,
                    campaign_id=campaign.id,
                    source_channel=source_channel,
                    source_medium=source_medium,
                    status=status_value,
                    current_stage=current_stage,
                    current_section=current_section,
                    current_handler=current_handler,
                    priority=priority,
                    lead_score=lead_score,
                    recommended_action=_next_action_hint(row),
                )
                db.add(lead)
                db.flush()
                created["leads"] += 1

                db.add(
                    LeadAssignment(
                        lead_id=lead.id,
                        customer_id=customer.id,
                        assignment_type="assign",
                        from_section=None,
                        to_section=current_section,
                        from_handler=None,
                        to_handler=current_handler,
                        reason="excel_preload",
                        trigger="dataset",
                        details={
                            "customer_id": customer_external_id,
                            "crm_channel": _str(row.get("CRM_Channel"), 64),
                            "crm_lead_type": _str(row.get("CRM_Lead_Type"), 64),
                        },
                        is_current=True,
                    )
                )
                created["lead_assignments"] += 1

                db.add(
                    LeadTimelineEvent(
                        lead_id=lead.id,
                        customer_id=customer.id,
                        event_type="lead_created",
                        event_source="sample_data.xlsx",
                        details={
                            "label_source_disposition": _str(row.get("Label_Source_Disposition"), 120),
                            "label_source_lead_status": _str(row.get("Label_Source_Lead_Status"), 120),
                            "label_basis": _str(row.get("Label_Basis"), 64),
                            "label_customer_validity": _as_int(row.get("LABEL_Customer_Validity"), default=0),
                        },
                    )
                )
                created["lead_timeline_events"] += 1

            existing_event = db.scalar(select(EngagementEvent.id).where(EngagementEvent.lead_id == lead.id))
            if existing_event is None:
                page_views = _as_int(row.get("Page_Views"), default=0)
                msg_engaged = _as_int(row.get("MSG_Engaged"), default=0)
                connected_calls = _as_int(row.get("CDR_Connected_Calls"), default=0)
                metric_value = max(1, page_views, msg_engaged, connected_calls)
                db.add(
                    EngagementEvent(
                        lead_id=lead.id,
                        customer_id=customer.id,
                        campaign_id=campaign.id,
                        channel=_slug(source_channel, 32),
                        metric_type="engaged",
                        metric_value=metric_value,
                        event_payload={
                            "crm_channel": _str(row.get("CRM_Channel"), 64),
                            "crm_department": _str(row.get("CRM_Department"), 64),
                            "crm_lead_source": _str(row.get("CRM_Lead_Source"), 64),
                            "crm_data_medium": _str(row.get("CRM_Data_Medium"), 64),
                            "msg_sent": _as_int(row.get("MSG_Sent"), default=0),
                            "msg_delivered": _as_int(row.get("MSG_Delivered"), default=0),
                            "msg_read": _as_int(row.get("MSG_Read"), default=0),
                            "msg_clicked": _as_int(row.get("MSG_Clicked"), default=0),
                        },
                    )
                )
                created["engagement_events"] += 1

            existing_web = db.scalar(select(WebsiteEvent.id).where(WebsiteEvent.lead_id == lead.id))
            if existing_web is None:
                step_name = _str(row.get("WEB_Step_Name"), 100) or "Unknown Step"
                db.add(
                    WebsiteEvent(
                        lead_id=lead.id,
                        customer_id=customer.id,
                        event_name=_slug(step_name, 64),
                        step_name=step_name,
                        step_number=_as_int(row.get("WEB_Step_Number"), default=1),
                        device_type=_str(row.get("WEB_Device_Type"), 32) or "unknown",
                        is_repeat_visitor=_as_bool(row.get("WEB_New_Vs_Repeat")),
                        event_payload={
                            "web_tracked": _as_int(row.get("WEB_Tracked"), default=0),
                            "visits": _as_int(row.get("Visits"), default=0),
                            "page_views": _as_int(row.get("Page_Views"), default=0),
                            "total_seconds_spent": _as_int(row.get("Total_Seconds_Spent"), default=0),
                            "plan_type": _str(row.get("WEB_Plan_Type"), 64),
                            "plan_variant": _str(row.get("WEB_Plan_Variant"), 64),
                            "quoted_price": _as_float(row.get("WEB_Quoted_Price"), default=0.0),
                            "coverage_amount": _as_float(row.get("WEB_Coverage_Amount"), default=0.0),
                        },
                    )
                )
                created["website_events"] += 1

            task_title = f"Handle lead {customer_external_id}"[:160]
            task = db.scalar(select(Task).where(Task.lead_id == lead.id, Task.title == task_title))
            if task is None:
                days_delta = max(1, min(14, int(round(_as_float(row.get("CRM_Days_Create_To_Update"), default=2.0)))))
                due_at = datetime.now(UTC) + timedelta(days=days_delta)
                task = Task(
                    lead_id=lead.id,
                    customer_id=customer.id,
                    title=task_title,
                    description=(
                        f"Disposition: {_str(row.get('Label_Source_Disposition'))}; "
                        f"LeadStatus: {_str(row.get('Label_Source_Lead_Status'))}; "
                        f"Basis: {_str(row.get('Label_Basis'))}"
                    )[:1000],
                    status="completed" if status_value == "converted" else "open",
                    priority=priority,
                    due_at=due_at,
                )
                db.add(task)
                db.flush()
                created["tasks"] += 1

            followup = db.scalar(select(Followup).where(Followup.task_id == task.id))
            if followup is None:
                followup = Followup(
                    task_id=task.id,
                    lead_id=lead.id,
                    customer_id=customer.id,
                    channel="call" if _as_int(row.get("CDR_Total_Calls"), default=0) > 0 else "email",
                    notes=(
                        f"Next action: {_next_action_hint(row)}; "
                        f"Top call status: {_str(row.get('CDR_Top_Call_Status'))}"
                    )[:1000],
                    status="completed" if status_value == "converted" else "pending",
                    scheduled_at=datetime.now(UTC) + timedelta(days=1),
                    completed_at=datetime.now(UTC) if status_value == "converted" else None,
                )
                db.add(followup)
                created["followups"] += 1

            total_calls = _as_int(row.get("CDR_Total_Calls"), default=0)
            if total_calls > 0:
                call = db.scalar(select(Call).where(Call.lead_id == lead.id, Call.customer_id == customer.id))
                if call is None:
                    avg_talk_sec = max(30, _as_int(row.get("CDR_Avg_Talk_Sec"), default=120))
                    ended = datetime.now(UTC) - timedelta(minutes=2)
                    started = ended - timedelta(seconds=avg_talk_sec)
                    db.add(
                        Call(
                            lead_id=lead.id,
                            customer_id=customer.id,
                            direction="inbound" if _as_int(row.get("CDR_Inbound_Calls"), default=0) > 0 else "outbound",
                            status="completed" if _as_int(row.get("CDR_Connected_Calls"), default=0) > 0 else "scheduled",
                            phone_number=None,
                            notes=f"Top status: {_str(row.get('CDR_Top_Call_Status'))}"[:1000],
                            started_at=started,
                            ended_at=ended,
                            duration_seconds=avg_talk_sec,
                        )
                    )
                    created["calls"] += 1

        db.commit()

        total_records = {
            "products": _count(db, Product),
            "campaigns": _count(db, Campaign),
            "customers": _count(db, Customer),
            "leads": _count(db, Lead),
            "lead_assignments": _count(db, LeadAssignment),
            "lead_timeline_events": _count(db, LeadTimelineEvent),
            "tasks": _count(db, Task),
            "followups": _count(db, Followup),
            "calls": _count(db, Call),
            "engagement_events": _count(db, EngagementEvent),
            "website_events": _count(db, WebsiteEvent),
        }

    print("real-excel-data-ready")
    print(f"dataset_path: {dataset_path}")
    print(f"dataset_columns: {len(columns)}")
    print(f"dataset_rows: {len(rows)}")
    print("new_rows:")
    for key, value in created.items():
        print(f"  {key}: {value}")
    print("total_rows:")
    for key, value in total_records.items():
        print(f"  {key}: {value}")


if __name__ == "__main__":
    main()
