#!/usr/bin/env python3
from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path
import sys
from typing import Any

from openpyxl import load_workbook
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer
from reportlab.graphics.shapes import Drawing, Line, Polygon, Rect, String

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from app.main import app

DOCS_DIR = ROOT / "docs"
MD_PATH = DOCS_DIR / "project_backend_handbook.md"
PDF_PATH = DOCS_DIR / "project_backend_handbook.pdf"
SAMPLE_XLSX_PATH = ROOT / "sample_data.xlsx"

HTTP_METHODS = ("get", "post", "put", "patch", "delete", "options", "head")


def _box(drawing: Drawing, x: float, y: float, w: float, h: float, text: str) -> None:
    drawing.add(Rect(x, y, w, h, strokeColor="#1F2937", fillColor="#E5E7EB", strokeWidth=1.2))
    drawing.add(String(x + 8, y + h / 2 - 4, text, fontName="Helvetica", fontSize=8.8, fillColor="#111827"))


def _arrow(drawing: Drawing, x1: float, y1: float, x2: float, y2: float) -> None:
    drawing.add(Line(x1, y1, x2, y2, strokeColor="#374151", strokeWidth=1.2))
    if x2 >= x1:
        head = Polygon(
            [x2, y2, x2 - 6, y2 + 3, x2 - 6, y2 - 3],
            strokeColor="#374151",
            fillColor="#374151",
        )
    else:
        head = Polygon(
            [x2, y2, x2 + 6, y2 + 3, x2 + 6, y2 - 3],
            strokeColor="#374151",
            fillColor="#374151",
        )
    drawing.add(head)


def architecture_diagram() -> Drawing:
    d = Drawing(520, 220)
    d.add(String(10, 205, "System Architecture", fontName="Helvetica-Bold", fontSize=11))

    _box(d, 12, 150, 86, 34, "Client")
    _box(d, 110, 150, 104, 34, "API Routes")
    _box(d, 226, 150, 104, 34, "Services")
    _box(d, 342, 150, 104, 34, "Repositories")
    _box(d, 458, 150, 56, 34, "Postgres")

    _arrow(d, 98, 167, 110, 167)
    _arrow(d, 214, 167, 226, 167)
    _arrow(d, 330, 167, 342, 167)
    _arrow(d, 446, 167, 458, 167)

    _box(d, 110, 90, 104, 34, "Pydantic Schemas")
    _box(d, 226, 90, 104, 34, "Workflow Rules")
    _box(d, 342, 90, 104, 34, "SQLAlchemy Models")
    _box(d, 458, 90, 56, 34, "Alembic")

    _arrow(d, 162, 150, 162, 124)
    _arrow(d, 278, 150, 278, 124)
    _arrow(d, 394, 150, 394, 124)
    _arrow(d, 486, 150, 486, 124)

    _box(d, 180, 25, 174, 38, "Sample Data + Import Pipeline")
    _arrow(d, 267, 90, 267, 63)

    d.add(String(12, 10, "Flow: request in -> domain orchestration -> persistent state -> typed response", fontName="Helvetica", fontSize=8.5))
    return d


def lifecycle_diagram() -> Drawing:
    d = Drawing(520, 235)
    d.add(String(10, 220, "Lead Lifecycle Workflow", fontName="Helvetica-Bold", fontSize=11))

    _box(d, 12, 165, 84, 32, "Lead Created")
    _box(d, 106, 165, 88, 32, "Assigned")
    _box(d, 204, 165, 88, 32, "Qualified")
    _box(d, 302, 165, 88, 32, "Decision")
    _box(d, 400, 165, 104, 32, "Action Outcome")

    _arrow(d, 96, 181, 106, 181)
    _arrow(d, 194, 181, 204, 181)
    _arrow(d, 292, 181, 302, 181)
    _arrow(d, 390, 181, 400, 181)

    _box(d, 120, 110, 120, 30, "Transfer")
    _box(d, 260, 110, 120, 30, "Follow-up Created")
    _box(d, 400, 110, 104, 30, "Status Update")

    _arrow(d, 160, 165, 180, 140)
    _arrow(d, 346, 165, 320, 140)
    _arrow(d, 452, 165, 452, 140)

    _box(d, 12, 48, 492, 42, "Timeline Events + Unified Details API (single payload for frontend lead workspace)")
    _arrow(d, 180, 110, 180, 90)
    _arrow(d, 320, 110, 260, 90)
    _arrow(d, 452, 110, 340, 90)

    d.add(String(12, 18, "Outcomes can move lead stage, set recommendation, and trigger new operational tasks.", fontName="Helvetica", fontSize=8.5))
    return d


def resolve_ref(ref: str, spec: dict[str, Any]) -> dict[str, Any]:
    parts = ref.strip("#/").split("/")
    node: Any = spec
    for part in parts:
        node = node.get(part, {})
    return node if isinstance(node, dict) else {}


def schema_keys(schema: dict[str, Any], spec: dict[str, Any]) -> list[str]:
    if not schema:
        return []
    if "$ref" in schema:
        return schema_keys(resolve_ref(schema["$ref"], spec), spec)
    if schema.get("type") == "array":
        item_keys = schema_keys(schema.get("items", {}), spec)
        return [f"items[{k}]" for k in item_keys] if item_keys else ["items"]
    if "allOf" in schema:
        merged: list[str] = []
        for sub in schema["allOf"]:
            merged.extend(schema_keys(sub, spec))
        return sorted(set(merged))
    props = schema.get("properties", {})
    if isinstance(props, dict) and props:
        return sorted(props.keys())
    return []


def response_summary(op: dict[str, Any], spec: dict[str, Any]) -> str:
    responses = op.get("responses", {})
    parts: list[str] = []
    for code in sorted(responses.keys()):
        info = responses.get(code, {})
        content = info.get("content", {}) if isinstance(info, dict) else {}
        schema = content.get("application/json", {}).get("schema", {})
        keys = schema_keys(schema, spec)
        if keys:
            preview = ", ".join(keys[:8])
            if len(keys) > 8:
                preview += ", ..."
            parts.append(f"{code}: {preview}")
        else:
            desc = info.get("description", "response") if isinstance(info, dict) else "response"
            parts.append(f"{code}: {desc}")
    return " | ".join(parts)


def request_summary(op: dict[str, Any], spec: dict[str, Any]) -> str:
    params = op.get("parameters", [])
    param_names = [f"{p.get('name')}({p.get('in')})" for p in params if isinstance(p, dict)]

    body_keys: list[str] = []
    req_body = op.get("requestBody", {})
    if isinstance(req_body, dict):
        content = req_body.get("content", {})
        schema = content.get("application/json", {}).get("schema", {})
        body_keys = schema_keys(schema, spec)

    parts: list[str] = []
    if param_names:
        parts.append("params=" + ", ".join(param_names))
    if body_keys:
        preview = ", ".join(body_keys[:10])
        if len(body_keys) > 10:
            preview += ", ..."
        parts.append("body=" + preview)
    return " ; ".join(parts) if parts else "none"


def load_excel_metadata(path: Path) -> tuple[list[str], list[tuple[str, str, str]]]:
    if not path.exists():
        return [], []

    workbook = load_workbook(path, read_only=True, data_only=True)

    headers: list[str] = []
    if "dataset" in workbook.sheetnames:
        dataset_sheet = workbook["dataset"]
        for row in dataset_sheet.iter_rows(min_row=1, max_row=1, values_only=True):
            headers = [str(value).strip() for value in row if value is not None and str(value).strip()]

    definitions: list[tuple[str, str, str]] = []
    if "data_dictionary" in workbook.sheetnames:
        dictionary_sheet = workbook["data_dictionary"]
        for row in dictionary_sheet.iter_rows(min_row=2, values_only=True):
            if row is None:
                continue
            column = str(row[0]).strip() if row[0] is not None else ""
            if not column:
                continue
            source = str(row[1]).strip() if len(row) > 1 and row[1] is not None else ""
            description = str(row[2]).strip() if len(row) > 2 and row[2] is not None else ""
            definitions.append((column, source, description))

    return headers, definitions


def build_markdown(spec: dict[str, Any]) -> str:
    now = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S UTC")
    paths = spec.get("paths", {})
    excel_headers, excel_definitions = load_excel_metadata(SAMPLE_XLSX_PATH)

    lines: list[str] = []
    lines.append("# Backend Project Handbook")
    lines.append("")
    lines.append(f"Generated: {now}")
    lines.append("")
    lines.append("## 1. Executive Summary")
    lines.append("")
    lines.append("This backend is a lead-centric sales intelligence system built on FastAPI, SQLAlchemy, Alembic, and PostgreSQL.")
    lines.append("The lead is the central business entity and is connected to customer, campaign, journey, engagement, call, task, follow-up, AI prediction, and decision modules.")
    lines.append("")
    lines.append("## 2. Architecture")
    lines.append("")
    lines.append("- API Layer: app/api/routes")
    lines.append("- Service Layer: app/services")
    lines.append("- Repository Layer: app/repositories")
    lines.append("- Persistence Layer: app/models + PostgreSQL")
    lines.append("- Schema Contracts: app/schemas")
    lines.append("- DB Evolution: Alembic migrations")
    lines.append("")
    lines.append("Request flow:")
    lines.append("")
    lines.append("```text")
    lines.append("Client -> Route -> Service -> Repository -> DB")
    lines.append("                         -> business rules -> response schema")
    lines.append("```")
    lines.append("")
    lines.append("Architecture Diagram (Mermaid):")
    lines.append("")
    lines.append("```mermaid")
    lines.append("flowchart LR")
    lines.append("  A[Client / Frontend] --> B[FastAPI Routes]")
    lines.append("  B --> C[Service Layer]")
    lines.append("  C --> D[Repository Layer]")
    lines.append("  D --> E[(PostgreSQL)]")
    lines.append("  B -. validation .-> F[Pydantic Schemas]")
    lines.append("  C -. rules .-> G[Lifecycle Orchestration]")
    lines.append("  D -. mapping .-> H[SQLAlchemy Models]")
    lines.append("  H -. evolve .-> I[Alembic Migrations]")
    lines.append("  J[Sample Data Loader] --> E")
    lines.append("```")
    lines.append("")
    lines.append("## 3. Feature Modules")
    lines.append("")
    features = [
        "Health + Dashboard",
        "Customers",
        "Leads",
        "Products",
        "Campaigns",
        "Engagement Channels",
        "Website Journey",
        "Calls",
        "Tasks + Follow-ups",
        "Data Import + Export",
        "ML Training + Model Catalog",
        "AI Predictions",
        "Decision Engine",
        "Analytics",
        "Reports",
        "Lead Lifecycle (assignment, transfer, outcomes, timeline, details)",
    ]
    for idx, feature in enumerate(features, start=1):
        lines.append(f"{idx}. {feature}")
    lines.append("")
    lines.append("## 4. Data Processing")
    lines.append("")
    lines.append("1. Source data is loaded from backend/sample_data.xlsx using scripts/preload_sample_data.py.")
    lines.append("2. The loader maps grouped records into products, campaigns, customers, leads, tasks, follow-ups, calls, engagement events, and website events.")
    lines.append("3. Live APIs use these persisted records, not synthetic smoke values.")
    lines.append("4. AI endpoints score/segment leads and store prediction logs.")
    lines.append("5. Decision endpoint combines predictions + rules into next action recommendations.")
    lines.append("6. Lifecycle module adds assignment/transfer, outcome tracking, and timeline events to create a connected business workflow.")
    lines.append("7. Unified lead details endpoint aggregates lead + customer + campaign + engagement + AI + operational activity for frontend rendering.")
    lines.append("")
    lines.append("### 4.1 Sample Data Columns (sample_data.xlsx)")
    lines.append("")
    if excel_headers:
        lines.append(f"Total columns in dataset sheet: {len(excel_headers)}")
        lines.append("")
        for idx, header in enumerate(excel_headers, start=1):
            lines.append(f"{idx}. {header}")
    else:
        lines.append("sample_data.xlsx dataset columns were not available at generation time.")
    lines.append("")
    lines.append("### 4.2 Column Meaning (data_dictionary sheet)")
    lines.append("")
    if excel_definitions:
        for column, source, description in excel_definitions:
            lines.append(f"- {column} | source={source or 'n/a'} | {description or 'n/a'}")
    else:
        lines.append("data_dictionary definitions were not available at generation time.")
    lines.append("")
    lines.append("### 4.3 Column To API Field Mapping")
    lines.append("")
    lines.append("- Customer_ID -> customers.external_customer_id -> returned in /api/customers and /api/leads/{id}/details.customer")
    lines.append("- CRM_Product_Code/CRM_Product_Name -> products.code/products.name -> returned in /api/products and details.campaign linkage")
    lines.append("- CRM_UTM_Campaign/CRM_Data_Source_Platform -> campaigns.code/name -> returned in /api/campaigns")
    lines.append("- CRM_Channel/CRM_Data_Medium/CRM_Source -> leads.source_channel/source_medium/current_handler")
    lines.append("- Label_Source_Lead_Status/Label_Source_Disposition -> leads.status/recommended_action")
    lines.append("- LABEL_Customer_Validity/CDR_Connect_Rate/Page_Views/MSG_Engaged -> leads.lead_score")
    lines.append("- CRM_Department -> leads.current_section and lifecycle assignment.to_section")
    lines.append("- CDR_* metrics -> calls and engagement event payloads")
    lines.append("- WEB_* metrics -> website_events fields and payload")
    lines.append("- Label_Basis/Disposition -> task descriptions + followup notes + timeline details")
    lines.append("")
    lines.append("## 5. How To Run In Detail")
    lines.append("")
    lines.append("1. Create backend-local virtual environment and install requirements.")
    lines.append("2. Configure .env with DATABASE_URL for PostgreSQL.")
    lines.append("3. Run migrations: ./.venv/bin/alembic upgrade head")
    lines.append("4. Preload real dataset: ./.venv/bin/python scripts/preload_sample_data.py --reset")
    lines.append("5. Start API server: ./.venv/bin/uvicorn app.main:app --host 0.0.0.0 --port 8000")
    lines.append("6. Open Swagger: /api/docs and OpenAPI: /api/openapi.json")
    lines.append("7. Run smoke verification: ./.venv/bin/python scripts/live_api_smoke.py")
    lines.append("8. Run tests: ./.venv/bin/pytest -q")
    lines.append("9. Lifecycle-only test: ./.venv/bin/pytest -q tests/test_lead_lifecycle.py")
    lines.append("")
    lines.append("## 6. Feature Workflows And Handling")
    lines.append("")
    lines.append("1. Lead intake: create lead -> initial timeline event -> assignment queue.")
    lines.append("2. Assignment/transfer: updates current owner and section; previous assignment marked non-current.")
    lines.append("3. AI scoring: prediction APIs persist score outputs per prediction type.")
    lines.append("4. Decisioning: next action recommendation combines prediction signal and business logic.")
    lines.append("5. Outcome handling: captures result, updates lead stage/status, optionally creates follow-up, and appends timeline events.")
    lines.append("6. Reporting: analytics/reports endpoints query aggregated module-level data for operational monitoring.")
    lines.append("")
    lines.append("Lifecycle Diagram (Mermaid):")
    lines.append("")
    lines.append("```mermaid")
    lines.append("flowchart LR")
    lines.append("  A[Lead Created] --> B[Assign or Transfer]")
    lines.append("  B --> C[AI + Decision]")
    lines.append("  C --> D[Action Taken]")
    lines.append("  D --> E[Outcome Recorded]")
    lines.append("  E --> F[Update Lead State]")
    lines.append("  E --> G[Create Follow-up]")
    lines.append("  B --> H[Timeline Event]")
    lines.append("  D --> H")
    lines.append("  E --> H")
    lines.append("  H --> I[Unified Lead Details API]")
    lines.append("```")
    lines.append("")
    lines.append("## 7. API Inventory And Response Contracts")
    lines.append("")
    lines.append(f"Total API paths in OpenAPI: {len(paths)}")
    lines.append("")

    for path in sorted(paths.keys()):
        methods = paths[path]
        for method in HTTP_METHODS:
            if method not in methods:
                continue
            op = methods[method]
            summary = op.get("summary", "")
            tag = ",".join(op.get("tags", []))
            req = request_summary(op, spec)
            resp = response_summary(op, spec)
            lines.append(f"### {method.upper()} {path}")
            if summary:
                lines.append(f"- Summary: {summary}")
            if tag:
                lines.append(f"- Tags: {tag}")
            lines.append(f"- Request: {req}")
            lines.append(f"- Responses: {resp}")
            lines.append("")

    lines.append("## 8. Lead Lifecycle Behavior")
    lines.append("")
    lines.append("1. New lead creation writes lead_created event to timeline.")
    lines.append("2. Assign/transfer updates lead ownership and section, and writes lifecycle events.")
    lines.append("3. Outcome recording can update lead status/stage and auto-create follow-ups.")
    lines.append("4. /leads/{id}/details returns a unified payload for frontend lead workspace rendering.")
    lines.append("")
    lines.append("## 9. Validation Snapshot")
    lines.append("")
    lines.append("- Test suite passing (including lifecycle tests).")
    lines.append("- Live smoke APIs passing with real sample data preloaded.")
    lines.append("- Lifecycle endpoints verified: timeline, details, assign, transfer, outcomes list/create.")
    lines.append("")

    return "\n".join(lines)


def markdown_to_pdf(md_text: str, pdf_path: Path) -> None:
    styles = getSampleStyleSheet()
    body = ParagraphStyle(
        "Body",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=9.5,
        leading=13,
        spaceAfter=4,
    )
    heading = ParagraphStyle(
        "Heading",
        parent=styles["Heading2"],
        fontName="Helvetica-Bold",
        fontSize=12,
        leading=15,
        spaceBefore=8,
        spaceAfter=6,
    )
    code = ParagraphStyle(
        "Code",
        parent=styles["Code"],
        fontName="Courier",
        fontSize=8.5,
        leading=11,
        spaceAfter=3,
    )
    subheading = ParagraphStyle(
        "SubHeading",
        parent=styles["Heading3"],
        fontName="Helvetica-Bold",
        fontSize=10.5,
        leading=13,
        spaceBefore=6,
        spaceAfter=4,
    )

    story = []
    mermaid_mode = False
    last_heading = ""
    for raw_line in md_text.splitlines():
        line = raw_line.strip()
        if not line:
            story.append(Spacer(1, 2.5 * mm))
            continue
        if line == "```mermaid":
            mermaid_mode = True
            continue
        if line == "```" and mermaid_mode:
            mermaid_mode = False
            continue
        if mermaid_mode:
            continue
        if line.startswith("### ") or line.startswith("## ") or line.startswith("# "):
            last_heading = line.lstrip("# ")
            story.append(Paragraph(last_heading, heading))
            if "Architecture" in last_heading:
                story.append(Spacer(1, 2 * mm))
                story.append(architecture_diagram())
                story.append(Spacer(1, 3 * mm))
            if "Feature Workflows" in last_heading:
                story.append(Spacer(1, 2 * mm))
                story.append(lifecycle_diagram())
                story.append(Spacer(1, 3 * mm))
            continue
        if line.startswith("```"):
            continue
        if "->" in line and "Client" in line:
            story.append(Paragraph(line.replace(" ", "&nbsp;"), code))
            continue
        if line.startswith("Architecture Diagram") or line.startswith("Lifecycle Diagram"):
            story.append(Paragraph(line, subheading))
            continue
        if line.startswith("- ") or line[0:2].isdigit() and line[1] == ".":
            story.append(Paragraph(line, body))
            continue
        story.append(Paragraph(line, body))

    doc = SimpleDocTemplate(
        str(pdf_path),
        pagesize=A4,
        leftMargin=15 * mm,
        rightMargin=15 * mm,
        topMargin=14 * mm,
        bottomMargin=14 * mm,
        title="Backend Project Handbook",
        author="AI Sales and Lead Intelligence Backend",
    )
    doc.build(story)


def main() -> None:
    DOCS_DIR.mkdir(parents=True, exist_ok=True)

    spec = app.openapi()
    md_text = build_markdown(spec)
    MD_PATH.write_text(md_text, encoding="utf-8")
    markdown_to_pdf(md_text, PDF_PATH)

    print(f"Markdown written: {MD_PATH}")
    print(f"PDF written: {PDF_PATH}")


if __name__ == "__main__":
    main()
