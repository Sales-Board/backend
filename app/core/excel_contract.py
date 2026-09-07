"""Canonical source-of-truth contract for backend Excel data.

The workbook is authoritative. These constants only name columns verified in
sample_data.xlsx; derived system fields are kept separate from source fields.
"""

from pathlib import Path

from openpyxl import load_workbook

ROOT_DIR = Path(__file__).resolve().parents[2]
DATASET_PATH = ROOT_DIR / "sample_data.xlsx"

LEAD_COLUMNS = (
    "Customer_ID",
    "CRM_Channel",
    "CRM_Department",
    "CRM_Lead_Type",
    "CRM_Lead_Source",
    "CRM_Data_Medium",
    "CRM_Source",
    "CRM_Data_Source_Platform",
    "CRM_Product_Code",
    "CRM_Product_Name",
    "CRM_UTM_Source",
    "CRM_UTM_Medium",
    "CRM_UTM_Campaign",
    "CRM_Lead_Create_Hour",
    "CRM_Lead_Create_DayOfWeek",
    "CRM_Lead_Create_Weekend",
    "CRM_Lead_Create_Month",
    "CRM_Days_Create_To_Update",
    "CRM_Contact_Count",
    "CRM_NonContact_Count",
    "CRM_No_Of_Attempts",
    "Label_Source_Disposition",
    "Label_Source_Lead_Status",
)

CUSTOMER_COLUMNS = (
    "Customer_ID",
    "CRM_Gender",
    "CRM_Age_Band",
    "CRM_Income_Band",
    "CRM_Occupation",
    "CRM_Education",
    "CRM_Tobacco_User",
    "CRM_NonResident_Flag",
    "CRM_Existing_Plan_Flag",
)

PRODUCT_COLUMNS = (
    "CRM_Product_Code",
    "CRM_Product_Name",
    "WEB_Plan_Type",
    "WEB_Plan_Variant",
    "WEB_Quoted_Price",
    "WEB_Coverage_Amount",
    "WEB_Payment_Frequency",
)

CAMPAIGN_COLUMNS = (
    "CRM_Channel",
    "CRM_Department",
    "CRM_Lead_Source",
    "CRM_Data_Medium",
    "CRM_Source",
    "CRM_Data_Source_Platform",
    "CRM_UTM_Source",
    "CRM_UTM_Medium",
    "CRM_UTM_Campaign",
    "MSG_Campaigns_Targeted",
)

MESSAGE_COLUMNS = (
    "MSG_Campaigns_Targeted",
    "MSG_Sent",
    "MSG_Delivered",
    "MSG_Read",
    "MSG_Clicked",
    "MSG_Replied",
    "MSG_Failed",
    "MSG_Engaged",
)

CALL_COLUMNS = (
    "CDR_Total_Calls",
    "CDR_Connected_Calls",
    "CDR_Connect_Rate",
    "CDR_Total_Talk_Sec",
    "CDR_Avg_Talk_Sec",
    "CDR_Max_Talk_Sec",
    "CDR_Distinct_Call_Days",
    "CDR_Distinct_Agents",
    "CDR_Callbacks_Scheduled",
    "CDR_Inbound_Calls",
    "CDR_Top_Call_Status",
    "CDR_Top_Q_Type",
    "CDR_First_Call_Delay_Days",
    "CDR_Call_Span_Days",
)

WEB_COLUMNS = (
    "WEB_Tracked",
    "Visits",
    "Page_Views",
    "Total_Seconds_Spent",
    "WEB_Step_Name",
    "WEB_Step_Number",
    "WEB_New_Vs_Repeat",
    "WEB_Plan_Type",
    "WEB_Plan_Variant",
    "WEB_Quoted_Price",
    "WEB_Coverage_Amount",
    "WEB_Payment_Frequency",
    "WEB_Last_Touch_Channel",
    "WEB_UTM_Source",
    "WEB_Device_Type",
)

EVENT_COLUMNS = (
    "ev_Lead_Creation",
    "ev_Lead_Submitted",
    "ev_Calculator_Start",
    "ev_Retrieve_Quote",
    "ev_Quote_Proceed",
    "ev_Personal_Details_Next",
    "ev_Plan_Details_Load",
    "ev_Make_Payment_Load",
    "ev_Pay_Now",
    "ev_Payment_Success",
    "ev_Payment_Failure",
    "ev_Proposal_Form_Load",
    "ev_Proposal_Form_Proceed",
    "ev_Generate_OTP",
    "ev_OTP_Submit",
    "ev_Brochure_Download",
    "ev_Connect_With_Expert",
    "ev_Page_Scroll_25",
    "ev_Page_Scroll_95",
)

TARGET_COLUMNS = ("LABEL_Customer_Validity",)
LEAKAGE_COLUMNS = frozenset(
    {
        "CRM_Has_Application_No",
        "CRM_Has_Payment_Flag",
        "ev_Payment_Success",
        "ev_Payment_Failure",
        "Label_Source_Disposition",
        "Label_Source_Lead_Status",
        "LABEL_Customer_Validity",
    }
)

PRE_OUTCOME_FEATURE_COLUMNS = tuple(
    column
    for column in dict.fromkeys(
        (*LEAD_COLUMNS, *CUSTOMER_COLUMNS, *PRODUCT_COLUMNS, *CAMPAIGN_COLUMNS, *MESSAGE_COLUMNS, *CALL_COLUMNS, *WEB_COLUMNS, *EVENT_COLUMNS)
    )
    if column not in LEAKAGE_COLUMNS
)


def workbook_columns(path: Path = DATASET_PATH) -> set[str]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook["dataset"]
    header = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
    return {str(value).strip() for value in header if value is not None and str(value).strip()}


def validate_contract(path: Path = DATASET_PATH) -> list[str]:
    available = workbook_columns(path)
    required = set(
        (*LEAD_COLUMNS, *CUSTOMER_COLUMNS, *PRODUCT_COLUMNS, *CAMPAIGN_COLUMNS, *MESSAGE_COLUMNS, *CALL_COLUMNS, *WEB_COLUMNS, *EVENT_COLUMNS, *TARGET_COLUMNS)
    )
    return sorted(required - available)


def source_features(excel_fields: dict | None) -> dict:
    values = excel_fields or {}
    return {column: values.get(column) for column in PRE_OUTCOME_FEATURE_COLUMNS if column in values}
